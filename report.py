# Charts and the Excel report (labels in Polish - the report is meant for a Polish client)
import matplotlib.pyplot as plt
import pandas as pd
from openpyxl.drawing.image import Image
from openpyxl.styles import Font


def daily_profile_chart(data, path):
    # average load and average price for each hour of the day
    avg = data.groupby(data.index.hour).mean()

    fig, ax1 = plt.subplots(figsize=(10, 5))
    ax1.bar(avg.index, avg["load_kw"], color="tab:green", alpha=0.6)
    ax1.set_xlabel("Godzina")
    ax1.set_ylabel("Średnia moc [kW]", color="tab:green")

    ax2 = ax1.twinx()
    ax2.plot(avg.index, avg["price"], color="tab:blue", marker="o")
    ax2.set_ylabel("Średnia cena RCE [zł/MWh]", color="tab:blue")

    plt.title("Profil dobowy klienta na tle cen RCE")
    plt.xticks(range(0, 24))
    plt.tight_layout()
    plt.savefig(path, dpi=100)
    plt.close()


def battery_chart(battery, path):
    fig, ax = plt.subplots(figsize=(10, 4))
    ax.bar(battery.index, battery["profit_pln"], color="tab:orange")
    ax.set_ylabel("Zysk [zł]")
    ax.set_title("Dzienny zysk z arbitrażu magazynu energii")
    plt.tight_layout()
    plt.savefig(path, dpi=100)
    plt.close()


def add_sum(ws, column, label="SUMA"):
    # Excel formula instead of a computed value - data can be edited later
    last = ws.max_row
    ws.cell(row=last + 1, column=1, value=label).font = Font(bold=True)
    cell = ws.cell(row=last + 1, column=column)
    letter = cell.column_letter
    cell.value = f"=SUM({letter}2:{letter}{last})"
    cell.font = Font(bold=True)


def save_excel(path, summary, dsr, battery, data, charts):
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        pd.Series(summary, name="Wartość").to_excel(writer, sheet_name="Podsumowanie")
        dsr.round(2).to_excel(writer, sheet_name="DSR")
        battery.round(2).to_excel(writer, sheet_name="Magazyn")
        data.to_excel(writer, sheet_name="Dane 15 min")

        ws = writer.sheets["Podsumowanie"]
        ws.column_dimensions["A"].width = 45
        ws.column_dimensions["B"].width = 25
        row = 2
        for c in charts:
            ws.add_image(Image(c), f"D{row}")
            row += 27

        add_sum(writer.sheets["DSR"], column=5)
        add_sum(writer.sheets["Magazyn"], column=6)

        for name in ["DSR", "Magazyn", "Dane 15 min"]:
            writer.sheets[name].column_dimensions["A"].width = 20
